"""Import generico da tabella (CSV/TSV o Excel .xlsx), riconoscendo le colonne dal nome.

Due forme di tabella:
- "a righe": una riga per articolo e taglia, con una colonna quantità
  (Articolo | Colore | Taglia | Quantità);
- "a colonne": una riga per articolo e una colonna per taglia con le quantità
  (Articolo | XS | S | M | L | …).
La descrizione si compone da articolo/descrizione + grammatura + colore, se ci sono.
Il risultato è solo una proposta: l'utente la controlla nel modulo prima di stampare.
"""
from __future__ import annotations

import csv
import io
import os
import re
import unicodedata
from typing import Dict, List, Optional, Tuple

from ..label.label import SIZES
from ..models import Articolo, Commessa, RigaTaglia
from .base import Importer, ImportNonDisponibile, MappaturaColonne, RisultatoImport

ESTENSIONI = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm")

# sinonimi delle intestazioni (già normalizzate: minuscolo, senza accenti e punteggiatura)
COMMESSA = ["commessa", "ordine", "n ordine", "numero ordine", "nr ordine", "riferimento", "rif", "cliente"]
DESCRIZIONE = ["descrizione", "descrizione articolo", "prodotto", "modello", "articolo", "capo", "item", "product"]
GRAMMATURA = ["grammatura", "peso", "gr", "gsm", "g m2", "gm2", "grammi"]
COLORE = ["colore", "color", "colour", "col"]
TAGLIA = ["taglia", "size", "tg", "misura"]
QUANTITA = ["quantita", "qta", "qty", "quantity", "q ta", "pezzi", "pz", "n capi", "capi", "copie", "numero capi"]

_XL = re.compile(r"^(\d)XL$")
_BIMBO = re.compile(r"^\d+(\s*[-/]\s*\d+)?\s*(A|ANNI|Y|YRS)$")


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", s.lower()).split())


def _taglia_da_intestazione(h) -> Optional[str]:
    """Nome di colonna che è una taglia (formato "a colonne"), altrimenti None."""
    t = " ".join(str(h or "").split()).upper()
    if not t:
        return None
    t = {"2XL": "XXL", "XXXL": "3XL"}.get(t, t)
    if t in SIZES or _XL.match(t) or _BIMBO.match(t) or t in ("UNICA", "TU", "TAGLIA UNICA"):
        return t
    return None


def _trova(intest: List[str], sinonimi: List[str]) -> Optional[int]:
    norm = [_norm(h) for h in intest]
    for s in sinonimi:              # in ordine di preferenza
        if s in norm:
            return norm.index(s)
    return None


def _quantita(v) -> Optional[int]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() and v >= 0 else None
    s = str(v).strip().replace(",", ".")
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f.is_integer() and f >= 0 else None


def _cella(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


# ---------------------------------------------------------------- lettura file
def _righe_csv(contenuto: bytes) -> List[List[str]]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            testo = contenuto.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    campione = testo[:4096]
    try:
        dialect = csv.Sniffer().sniff(campione, delimiters=";,\t|")
    except csv.Error:
        class dialect(csv.excel):
            delimiter = ";" if campione.count(";") > campione.count(",") else ","
    return [r for r in csv.reader(io.StringIO(testo), dialect)]


def _righe_xlsx(contenuto: bytes) -> List[list]:
    try:
        from openpyxl import load_workbook
    except ImportError:                                   # pragma: no cover
        raise ImportNonDisponibile("Per leggere i file Excel serve il pacchetto openpyxl.")
    try:
        wb = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
    except Exception:
        raise ImportNonDisponibile("Il file Excel non si apre: è danneggiato o protetto da password?")
    # il foglio con più celle piene (di solito l'unico)
    migliore, righe_migliori = 0, []
    for ws in wb.worksheets:
        righe = [list(r) for r in ws.iter_rows(values_only=True)]
        piene = sum(1 for r in righe for v in r if v not in (None, ""))
        if piene > migliore:
            migliore, righe_migliori = piene, righe
    wb.close()
    return righe_migliori


# ---------------------------------------------------------------- importer
class ImporterTabella(Importer):
    nome = "Tabella (CSV/Excel)"

    def riconosce(self, nome_file: str, contenuto: bytes) -> bool:
        return nome_file.lower().endswith(ESTENSIONI)

    def _righe(self, nome_file: str, contenuto: bytes) -> List[list]:
        if nome_file.lower().endswith((".xlsx", ".xlsm")):
            return _righe_xlsx(contenuto)
        return _righe_csv(contenuto)

    def colonne(self, contenuto: bytes, nome_file: str = "ordine.csv") -> List[str]:
        righe = self._righe(nome_file, contenuto)
        return [_cella(h) for h in righe[0]] if righe else []

    @staticmethod
    def _intestazione(righe: List[list]) -> Tuple[int, MappaturaColonne]:
        """Prima riga (entro le prime 20) che ha una colonna articolo e o una colonna
        quantità o almeno due colonne taglia."""
        for i, r in enumerate(righe[:20]):
            intest = [_cella(h) for h in r]
            d = _trova(intest, DESCRIZIONE)
            if d is None:
                continue
            m = MappaturaColonne(descrizione=intest[d])
            c = _trova(intest, COMMESSA)
            m.commessa = intest[c] if c is not None else None
            q = _trova(intest, QUANTITA)
            t = _trova(intest, TAGLIA)
            if q is not None:
                m.copie = intest[q]
                m.taglia = intest[t] if t is not None else None
                return i, m
            col_t = {intest[k]: tg for k, h in enumerate(intest) if (tg := _taglia_da_intestazione(h))}
            if len(col_t) >= 2:
                m.colonne_taglia = col_t
                return i, m
        raise ImportNonDisponibile("")

    def importa(self, contenuto: bytes, mappatura: Optional[MappaturaColonne] = None,
                nome_file: str = "ordine.csv") -> RisultatoImport:
        righe = self._righe(nome_file, contenuto)
        if not righe:
            raise ImportNonDisponibile("Il file è vuoto.")
        try:
            h, m = self._intestazione(righe)
        except ImportNonDisponibile:
            trovate = [c for c in (_cella(x) for x in righe[0]) if c][:12]
            raise ImportNonDisponibile(
                "Non riconosco le colonne del file. Servono una colonna Articolo (o Descrizione/Prodotto) e "
                "una colonna Quantità con la Taglia, oppure una colonna per ogni taglia (XS, S, M…). "
                f"Colonne trovate: {', '.join(trovate) or 'nessuna'}.")
        intest = [_cella(x) for x in righe[h]]
        idx = {c: k for k, c in enumerate(intest)}
        g = _trova(intest, GRAMMATURA)
        co = _trova(intest, COLORE)
        d = idx[m.descrizione]

        def val(r, k):
            return _cella(r[k]) if k is not None and k < len(r) else ""

        avvisi: List[str] = []
        commesse: List[str] = []
        articoli: Dict[Tuple[str, str], Dict[str, int]] = {}      # (commessa, descrizione) -> taglia -> capi
        for n, r in enumerate(righe[h + 1:], start=h + 2):
            if not any(_cella(v) for v in r):
                continue
            descr = val(r, d)
            gr = val(r, g)
            if gr and re.fullmatch(r"\d+([.,]\d+)?", gr):
                gr += " g"
            col = val(r, co)
            parti = [descr] + [p for p in (gr, col) if p and p.lower() not in descr.lower()]
            descr = " ".join(p for p in parti if p)
            nome = val(r, idx.get(m.commessa)) if m.commessa else ""
            if m.copie:
                q = _quantita(r[idx[m.copie]] if idx[m.copie] < len(r) else None)
                tg = val(r, idx.get(m.taglia)) if m.taglia else "Unica"
                if q is None:
                    avvisi.append(f"Riga {n}: quantità non valida, riga saltata.")
                    continue
                coppie = [(tg or "Unica", q)]
            else:
                coppie = []
                for c, tg in m.colonne_taglia.items():
                    k = idx[c]
                    q = _quantita(r[k] if k < len(r) else None)
                    if q is None:
                        avvisi.append(f"Riga {n}, colonna {c}: quantità non valida, saltata.")
                        continue
                    coppie.append((tg, q))
            coppie = [(t, q) for t, q in coppie if q > 0]
            if not coppie:
                continue
            if not descr:
                avvisi.append(f"Riga {n}: articolo senza descrizione.")
            if nome not in commesse:
                commesse.append(nome)
            art = articoli.setdefault((nome, descr), {})
            for t, q in coppie:
                art[t] = art.get(t, 0) + q

        if not articoli:
            raise ImportNonDisponibile("Nel file non ci sono righe con quantità maggiori di zero.")
        nome = commesse[0]
        if len(commesse) > 1:
            altre = ", ".join(c or "(senza nome)" for c in commesse[1:6])
            avvisi.insert(0, f"Il file contiene {len(commesse)} commesse: importata solo «{nome or '(senza nome)'}». "
                             f"Le altre ({altre}) vanno importate separatamente.")
        if not nome:
            nome = os.path.splitext(os.path.basename(nome_file))[0].replace("_", " ").strip()
        arts = []
        for (cn, descr), tg in articoli.items():
            if cn != commesse[0]:
                continue
            arts.append(Articolo(descrizione=descr[:300],
                                 righe=[RigaTaglia(taglia=t[:40], copie=min(q, 9999)) for t, q in tg.items()]))
        return RisultatoImport(commessa=Commessa(nome=nome[:200], articoli=arts), colonne=intest,
                               mappatura=m, avvisi=avvisi)
